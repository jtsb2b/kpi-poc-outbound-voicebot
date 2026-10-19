import io
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from kpi_voicebot.config import Settings
from kpi_voicebot.store import Store
from kpi_voicebot.worker import Worker


def call_body(**overrides: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "campaign": "accident_survey",
        "target_name": "สมชาย ใจดี",
        "phone": "081-234-5678",
        "mode": "simulation",
        "source": {"type": "profile", "profile_id": "survey-surveyor"},
        "scenario_id": "survey_contextual",
    }
    body.update(overrides)
    return body


async def create(client: AsyncClient, key: str = "key-00000001", **overrides: Any) -> Any:
    return await client.post("/api/v1/calls", json=call_body(**overrides), headers={"Idempotency-Key": key})


async def test_auth_required(client: AsyncClient) -> None:
    r = await client.get("/api/v1/calls", headers={"Authorization": "Bearer wrong-token-value"})
    assert r.status_code == 401
    assert (await client.get("/health")).status_code == 200
    assert (await client.get("/ready")).json()["mongo"] is True


async def test_campaigns_list_profiles_scenarios_and_pending_wording(client: AsyncClient) -> None:
    data = (await client.get("/api/v1/campaigns")).json()
    pa = next(c for c in data if c["id"] == "pa_confirmation")
    assert pa["persona_name"] == "น้องเคพีไอ"
    assert "free_look" in pa["pending_approval"]
    assert {p["profile_id"] for p in pa["profiles"]} == {"pa-email", "pa-post"}


async def test_simulated_call_end_to_end(client: AsyncClient, store: Store, settings: Settings) -> None:
    r = await create(client)
    assert r.status_code == 201
    call = r.json()
    assert call["phone"]["e164"] == "+66812345678" and call["phone"]["sip_destination"] == "0812345678"
    assert call["lifecycle"] == "queued"

    assert await Worker(store, settings, "w1").process_once()
    detail = (await client.get(f"/api/v1/calls/{call['call_id']}")).json()
    assert detail["lifecycle"] == "ended" and detail["dial_result"] == "answered"
    assert detail["outcome"] == "survey_completed"
    assert [a["value"] for a in detail["result"]["survey_answers"]] == [4, 3, 2, 5, 8]
    assert detail["transcript"][0]["speaker"] == "bot"
    assert [e["type"] for e in detail["events"]][:4] == ["queued", "dispatched", "dialing", "ringing"]
    assert detail["active"] is False


async def test_idempotency_and_single_active_call(client: AsyncClient) -> None:
    first = await create(client, "key-aaaaaaaa")
    again = await create(client, "key-aaaaaaaa")
    assert again.status_code == 200 and again.json()["call_id"] == first.json()["call_id"]
    changed = await create(client, "key-aaaaaaaa", target_name="คนอื่น")
    assert changed.status_code == 409
    second = await create(client, "key-bbbbbbbb")
    assert second.status_code == 409 and "กำลังดำเนินการ" in second.json()["detail"]


async def test_cancel_queued_call_never_dials(client: AsyncClient, store: Store, settings: Settings) -> None:
    call_id = (await create(client)).json()["call_id"]
    canceled = (await client.post(f"/api/v1/calls/{call_id}/cancel")).json()
    assert canceled["outcome"] == "canceled" and canceled["dial_result"] == "canceled"
    assert not await Worker(store, settings, "w1").process_once()
    assert (await create(client, "key-cccccccc")).status_code == 201  # lock released


async def test_cancel_during_conversation_keeps_partial_result(
    client: AsyncClient, store: Store, settings: Settings
) -> None:
    call_id = (await create(client)).json()["call_id"]
    worker = Worker(store, settings, "w1")
    original = store.add_turn
    turns = 0

    async def add_turn_then_cancel(cid: str, turn: dict[str, Any]) -> None:
        nonlocal turns
        await original(cid, turn)
        turns += 1
        if turns == 8:
            await store.request_cancel(cid)

    store.add_turn = add_turn_then_cancel  # type: ignore[method-assign]
    await worker.process_once()
    detail = (await client.get(f"/api/v1/calls/{call_id}")).json()
    assert detail["outcome"] == "canceled" and detail["end_reason"] == "operator_canceled"
    assert detail["result"]["identity_result"] == "confirmed"


async def test_rejections(client: AsyncClient) -> None:
    assert (await create(client, phone="+1 650 253 0000")).json()["detail"] == "ยังไม่รองรับการโทรต่างประเทศ"
    assert (await create(client, phone="1234")).status_code == 422
    assert (await create(client, mode="live")).status_code == 409
    wrong = await create(client, source={"type": "profile", "profile_id": "pa-email"})
    assert wrong.status_code == 422
    assert (await create(client, scenario_id=None)).status_code == 422


async def test_notice_objection_discards_transcript(client: AsyncClient, store: Store, settings: Settings) -> None:
    call_id = (
        await create(
            client,
            campaign="pa_confirmation",
            scenario_id="pa_objection",
            source={"type": "profile", "profile_id": "pa-email"},
        )
    ).json()["call_id"]
    await Worker(store, settings, "w1").process_once()
    detail = (await client.get(f"/api/v1/calls/{call_id}")).json()
    assert detail["outcome"] == "notice_objection"
    assert detail["transcript"] == [] and detail["transcript_discarded"] is True


async def test_custom_script(client: AsyncClient, store: Store, settings: Settings) -> None:
    call_id = (await create(client, scenario_id=None, custom_script=["ไม่ใช่ค่ะ ผิดคน"])).json()["call_id"]
    await Worker(store, settings, "w1").process_once()
    assert (await client.get(f"/api/v1/calls/{call_id}")).json()["outcome"] == "wrong_person"


async def test_lease_expiry_is_reconciled_not_redialled(store: Store, settings: Settings, client: AsyncClient) -> None:
    call_id = (await create(client)).json()["call_id"]
    claimed = await store.claim_next("dead-worker", lease_seconds=-1)
    assert claimed is not None
    assert await store.mark_dial_attempted(call_id, "dead-worker")
    assert await store.reconcile_expired() == [call_id]
    doc = await store.get_call(call_id)
    assert doc is not None and doc["outcome"] == "technical_failure" and doc["lifecycle"] == "ended"
    assert not await Worker(store, settings, "w2").process_once()


def _workbook(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_excel_import_select_row_and_export(client: AsyncClient, store: Store, settings: Settings) -> None:
    content = _workbook(
        [
            ["record_id", "case_id", "accident_date", "surveyor_attended", "customer_name", "phone"],
            ["R1", "CLM-1", "28/09/2569", "N", "สมหญิง ทดสอบ", "0812345678"],
            ["R2", "CLM-2", "not a date", "maybe", None, None],
            ["R3", '=HYPERLINK("http://example.com")', "2026-09-28", "", None, None],
        ]
    )
    r = await client.post(
        "/api/v1/imports",
        data={"campaign": "accident_survey"},
        files={"file": ("cases.xlsx", content, "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["valid_count"] == 1
    assert data["rows"][0]["record"]["accident_date"] == "2026-09-28"
    assert data["rows"][0]["record"]["surveyor_attended"] is False
    assert len(data["rows"][1]["errors"]) == 2
    assert data["rows"][2]["record"] is None  # formula cells are never evaluated

    created = await create(
        client,
        source={"type": "import", "import_id": data["import_id"], "row_number": 2},
        custom_script=["ใช่ค่ะ", "ได้ค่ะ", "ค่ะ", "5", "4", "10", "@SUM(A1)"],
        scenario_id=None,
    )
    assert created.status_code == 201, created.text
    assert any("ชื่อในไฟล์" in w for w in created.json()["warnings"])
    await Worker(store, settings, "w1").process_once()

    export = await client.get("/api/v1/exports/calls.xlsx", params={"call_id": created.json()["call_id"]})
    assert export.status_code == 200
    wb = load_workbook(io.BytesIO(export.content))
    assert set(wb.sheetnames) >= {"Summary", "Survey", "FollowUps", "Transcript"}
    summary = wb["Summary"]
    header = [c.value for c in summary[1]]
    row = {h: c.value for h, c in zip(header, summary[2], strict=True)}
    assert row["phone_entered"] == "081-234-5678" and row["outcome"] == "survey_completed"
    comment = next(r for r in wb["Survey"].iter_rows(values_only=True) if r[1] == "final_comment")
    assert comment[-1] == "'@SUM(A1)"
    for sheet in wb.worksheets:
        for r in sheet.iter_rows():
            for cell in r:
                assert cell.data_type != "f"


async def test_import_rejects_missing_columns(client: AsyncClient) -> None:
    content = _workbook([["record_id"], ["R1"]])
    r = await client.post(
        "/api/v1/imports",
        data={"campaign": "pa_confirmation"},
        files={"file": ("x.xlsx", content, "application/octet-stream")},
    )
    assert r.status_code == 422 and "plan_name" in r.json()["detail"]


async def test_template_download(client: AsyncClient) -> None:
    r = await client.get("/api/v1/templates/pa_confirmation.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb["records"][1]][:2] == ["record_id", "plan_name"]


async def test_worker_stops_claiming_calls_once_told_to_stop(store: Store, settings: Settings) -> None:
    import asyncio

    stop = asyncio.Event()
    stop.set()
    await asyncio.wait_for(Worker(store, settings, "w1").run_forever(stop), 1)  # returns instead of polling forever


def test_live_calls_keep_to_the_oic_calling_hours(settings: Settings) -> None:
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from kpi_voicebot.service import within_calling_hours

    bkk = ZoneInfo("Asia/Bangkok")
    assert within_calling_hours(settings, datetime(2026, 10, 19, 8, 30, tzinfo=bkk))  # Monday 08:30
    assert not within_calling_hours(settings, datetime(2026, 10, 19, 8, 29, tzinfo=bkk))
    assert not within_calling_hours(settings, datetime(2026, 10, 19, 19, 0, tzinfo=bkk))
    assert not within_calling_hours(settings, datetime(2026, 10, 18, 10, 0, tzinfo=bkk))  # Sunday

"""Excel import of campaign records and export of call results.

Excel is an exchange format only; MongoDB stays the source of truth. Imported formulas
are never evaluated (cached values only) and every exported text cell that could be
read as a formula is neutralised.
"""

import io
import re
from datetime import date, datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from kpi_voicebot.campaigns import load_profiles, validate_case
from kpi_voicebot.models import CampaignId

MAX_ROWS = 500

COLUMNS: dict[CampaignId, list[tuple[str, str, bool]]] = {
    # (column, description, required)
    CampaignId.PA_CONFIRMATION: [
        ("record_id", "รหัสรายการ (ไม่ซ้ำ)", True),
        ("plan_name", "ชื่อแผน เช่น พีเอ เต็มใจ", True),
        ("plan_code", "รหัสแผน", False),
        ("death_benefit_thb", "วงเงินคุ้มครองหลัก กรณีเสียชีวิตจากอุบัติเหตุ (บาท)", True),
        ("medical_per_accident_thb", "ค่ารักษาพยาบาลต่ออุบัติเหตุ (บาท) ถ้ามี", False),
        ("daily_hospital_thb", "เงินชดเชยรายได้ต่อวัน (บาท) ถ้ามี", False),
        ("daily_hospital_max_days", "จำนวนวันสูงสุดของเงินชดเชย", False),
        ("motorcycle_benefit_thb", "ความคุ้มครองกรณีรถจักรยานยนต์ (บาท) สำหรับตอบคำถาม", False),
        ("public_transport_benefit_thb", "ความคุ้มครองยานพาหนะสาธารณะ/รถยนต์ส่วนบุคคล (บาท)", False),
        ("premium_thb", "เบี้ยประกันภัย (บาท)", True),
        ("billing_period", "ปี หรือ เดือน", True),
        ("effective_date", "วันเริ่มคุ้มครอง (YYYY-MM-DD หรือ DD/MM/YYYY)", True),
        ("effective_time", "เวลาเริ่มคุ้มครอง เช่น 00.01 น.", False),
        ("beneficiary", "ผู้รับผลประโยชน์", True),
        ("payment_channel", "ช่องทางชำระเงิน", True),
        ("payment_status", "สถานะการชำระเงิน", False),
        ("delivery_channel", "อีเมล หรือ ไปรษณีย์", True),
        ("delivery_target", "อีเมลหรือที่อยู่จัดส่ง", True),
        ("delivery_period_text", "ระยะเวลาจัดส่ง เช่น 7 ถึง 14 วันทำการ", False),
        ("verify_birth_date", "วันเกิด (ใช้เมื่อเปิดการยืนยันตัวตนขั้นที่สอง)", False),
        ("customer_name", "ชื่อลูกค้าในไฟล์ (ใช้ตรวจสอบเท่านั้น ระบบใช้ชื่อจากฟอร์ม)", False),
        ("phone", "เบอร์ในไฟล์ (ใช้ตรวจสอบเท่านั้น ระบบโทรตามเบอร์ในฟอร์ม)", False),
    ],
    CampaignId.ACCIDENT_SURVEY: [
        ("record_id", "รหัสรายการ (ไม่ซ้ำ)", True),
        ("case_id", "เลขที่เคส/เคลม", True),
        ("accident_date", "วันที่เกิดเหตุ (YYYY-MM-DD หรือ DD/MM/YYYY)", True),
        ("surveyor_attended", "มีพนักงานสำรวจภัยไปที่เกิดเหตุ: Y / N / เว้นว่างถ้าไม่ทราบ", False),
        ("verify_birth_date", "วันเกิด (ใช้เมื่อเปิดการยืนยันตัวตนขั้นที่สอง)", False),
        ("customer_name", "ชื่อลูกค้าในไฟล์ (ใช้ตรวจสอบเท่านั้น)", False),
        ("phone", "เบอร์ในไฟล์ (ใช้ตรวจสอบเท่านั้น)", False),
    ],
}
_DATE_COLUMNS = {"effective_date", "accident_date", "verify_birth_date"}
_NUMBER_COLUMNS = {
    "death_benefit_thb",
    "medical_per_accident_thb",
    "daily_hospital_thb",
    "daily_hospital_max_days",
    "motorcycle_benefit_thb",
    "public_transport_benefit_thb",
    "premium_thb",
}
_INFO_COLUMNS = {"customer_name", "phone"}
_TRUE = {"y", "yes", "true", "1", "ใช่", "มี"}
_FALSE = {"n", "no", "false", "0", "ไม่ใช่", "ไม่มี"}
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


class ImportError_(ValueError):
    pass


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        y, mo, d = (int(g) for g in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
        if not m:
            raise ValueError("รูปแบบวันที่ต้องเป็น YYYY-MM-DD หรือ DD/MM/YYYY")
        d, mo, y = (int(g) for g in m.groups())
    if y > 2400:  # Buddhist era
        y -= 543
    return date(y, mo, d)


def _parse_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "").strip())


def _parse_bool(value: Any) -> bool | None:
    if value in (None, ""):
        return None
    text = str(value).strip().lower()
    if text in _TRUE:
        return True
    if text in _FALSE:
        return False
    raise ValueError("ต้องเป็น Y, N หรือเว้นว่าง")


def parse_workbook(campaign: CampaignId, content: bytes) -> list[dict[str, Any]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("ไม่สามารถอ่านไฟล์ Excel ได้ (.xlsx เท่านั้น)") from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise ImportError_("ไม่พบหัวตาราง")
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    spec = COLUMNS[campaign]
    missing = [c for c, _, required in spec if required and c not in headers]
    if missing:
        raise ImportError_(f"ไม่พบคอลัมน์ที่จำเป็น: {', '.join(missing)}")
    known = {c for c, _, _ in spec}

    parsed: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=2):
        if all(v in (None, "") for v in values):
            continue
        if len(parsed) >= MAX_ROWS:
            raise ImportError_(f"ไฟล์มีมากกว่า {MAX_ROWS} แถว")
        raw = {h: v for h, v in zip(headers, values, strict=False) if h in known}
        errors: list[str] = []
        warnings: list[str] = []
        case: dict[str, Any] = {}
        info: dict[str, str] = {}
        for column, value in raw.items():
            try:
                if column in _INFO_COLUMNS:
                    if value not in (None, ""):
                        if column == "phone" and isinstance(value, (int, float)):
                            warnings.append("เบอร์โทรในไฟล์ถูกเก็บเป็นตัวเลข เลข 0 นำหน้าอาจหายไป")
                        info[column] = str(int(value)) if isinstance(value, float) else str(value).strip()
                elif column in _DATE_COLUMNS:
                    case[column] = _parse_date(value)
                elif column in _NUMBER_COLUMNS:
                    case[column] = _parse_number(value)
                elif column == "surveyor_attended":
                    case[column] = _parse_bool(value)
                elif value not in (None, ""):
                    case[column] = str(value).strip()
            except ValueError as exc:
                errors.append(f"{column}: {exc}")
        case = {k: v for k, v in case.items() if v is not None or k == "surveyor_attended"}
        case["data_label"] = "IMPORT"
        record: dict[str, Any] | None = None
        if not errors:
            try:
                record = validate_case(campaign, case).model_dump(mode="json")
            except ValidationError as exc:
                errors += [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()]
        parsed.append(
            {
                "row_number": row_number,
                "record_id": str(raw.get("record_id") or ""),
                "record": record,
                "file_info": info,
                "errors": errors,
                "warnings": warnings,
            }
        )
    if not parsed:
        raise ImportError_("ไม่พบข้อมูลในไฟล์")
    return parsed


# ---- export ----


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIX):
        return "'" + value
    return value


def _write(ws: Worksheet, header: list[str], rows: list[list[Any]], text_columns: set[int] | None = None) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="13315F")
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if text_columns and cell.column - 1 in text_columns:
                cell.number_format = "@"
    for idx, title in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(12, min(48, len(title) + 6))
    ws.freeze_panes = "A2"


def _ts(value: Any, tz: ZoneInfo) -> str:
    if not isinstance(value, datetime):
        return ""
    from kpi_voicebot.store import as_utc

    return as_utc(value).astimezone(tz).strftime("%Y-%m-%d %H:%M:%S")


def export_calls(calls: list[dict[str, Any]], transcripts: list[dict[str, Any]], timezone: str) -> bytes:
    tz = ZoneInfo(timezone)
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = "Summary"
    summary_rows = []
    policy_rows: list[list[Any]] = []
    survey_rows: list[list[Any]] = []
    task_rows: list[list[Any]] = []
    for c in calls:
        result = c.get("result") or {}
        tasks = result.get("tasks", [])
        callback = result.get("callback") or {}
        summary_rows.append(
            [
                c["_id"],
                c["campaign"],
                c.get("mode"),
                c.get("target_name"),
                c["phone"].get("entered"),
                c["phone"].get("e164") or c["phone"].get("display"),
                c.get("data_label"),
                _source_label(c),
                _ts(c.get("created_at"), tz),
                _ts(c.get("answered_at"), tz),
                _ts(c.get("ended_at"), tz),
                c.get("lifecycle"),
                c.get("dial_result"),
                c.get("outcome"),
                c.get("end_reason"),
                result.get("identity_result"),
                "ใช่" if callback.get("requested") else "",
                callback.get("preferred_time_text"),
                len(tasks),
                "; ".join(sorted({t["type"] for t in tasks})),
                "ใช่" if c.get("transcript_discarded") else "",
            ]
        )
        for f in result.get("policy_fields", []):
            policy_rows.append(
                [
                    c["_id"],
                    f["field"],
                    f["label"],
                    f["expected_value"],
                    f["status"],
                    f["delivery"],
                    f.get("customer_statement"),
                ]
            )
        for a in result.get("survey_answers", []):
            survey_rows.append(
                [
                    c["_id"],
                    a["question"],
                    a["label"],
                    a.get("value"),
                    a["status"],
                    ", ".join(str(v) for v in a.get("changed_from", [])),
                    a.get("raw_utterance"),
                    a.get("feedback_raw"),
                ]
            )
        if c["campaign"] == CampaignId.ACCIDENT_SURVEY and (
            result.get("final_comment_raw") or result.get("final_comment_tags")
        ):
            survey_rows.append(
                [
                    c["_id"],
                    "final_comment",
                    "ความคิดเห็นเพิ่มเติม (ตามคำพูดลูกค้า)",
                    None,
                    "answered",
                    "",
                    None,
                    result.get("final_comment_raw"),
                ]
            )
        for t in tasks:
            task_rows.append([c["_id"], t["type"], t["reason"], t.get("field"), t.get("detail_raw")])

    _write(
        summary,
        [
            "call_id",
            "campaign",
            "mode",
            "target_name",
            "phone_entered",
            "phone_normalized",
            "data_label",
            "source",
            "created_at",
            "answered_at",
            "ended_at",
            "lifecycle",
            "dial_result",
            "outcome",
            "end_reason",
            "identity_result",
            "callback_requested",
            "callback_time_text",
            "follow_up_count",
            "follow_up_types",
            "transcript_discarded",
        ],
        summary_rows,
        text_columns={4, 5},
    )
    if policy_rows:
        _write(
            wb.create_sheet("PolicyConfirmation"),
            ["call_id", "field", "label", "expected_value", "status", "delivery", "customer_statement"],
            policy_rows,
        )
    if survey_rows:
        _write(
            wb.create_sheet("Survey"),
            [
                "call_id",
                "question",
                "label",
                "score",
                "status",
                "changed_from",
                "raw_utterance",
                "feedback_raw (verbatim)",
            ],
            survey_rows,
        )
    _write(wb.create_sheet("FollowUps"), ["call_id", "type", "reason", "field", "customer_words"], task_rows)
    _write(
        wb.create_sheet("Transcript"),
        ["call_id", "seq", "time", "speaker", "spoken_text", "generated_text", "interrupted"],
        [
            [
                t["call_id"],
                t["seq"],
                _ts(t.get("at"), tz),
                t["speaker"],
                t["text"],
                t.get("generated_text"),
                "ใช่" if t.get("interrupted") else "",
            ]
            for t in transcripts
        ],
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _source_label(call: dict[str, Any]) -> str:
    source = call.get("source") or {}
    if source.get("type") == "import":
        return f"Excel row {source.get('row_number')} ({source.get('record_id')})"
    return f"profile {source.get('profile_id')}"


def template_workbook(campaign: CampaignId) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "records"
    spec = COLUMNS[campaign]
    example = next(p for p in load_profiles() if p.campaign == campaign).case
    row = [example.get(c) for c, _, _ in spec]
    row = [v.isoformat() if isinstance(v, date) else v for v in row]
    _write(ws, [c for c, _, _ in spec], [row])
    for idx, (column, _, _) in enumerate(spec):
        if column in {"phone", "record_id", "case_id"}:
            for cell in ws.iter_cols(min_col=idx + 1, max_col=idx + 1, min_row=2):
                for c in cell:
                    c.number_format = "@"
    guide = wb.create_sheet("คำอธิบาย")
    _write(guide, ["column", "description", "required"], [[c, d, "จำเป็น" if r else ""] for c, d, r in spec])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
